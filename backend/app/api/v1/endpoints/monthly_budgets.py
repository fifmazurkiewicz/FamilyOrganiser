import uuid
from io import BytesIO

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.api.deps import get_current_user
from app.db.base import get_db
from app.models.user import User
from app.services.monthly_budget import MonthlyBudgetService
from app.schemas.monthly_budget import (
    MonthlyBudgetCreate, MonthlyBudgetResponse,
    BudgetEntryCreate, BudgetEntryUpdate, BudgetEntryResponse,
    BudgetSummaryResponse,
)

router = APIRouter(prefix="/monthly-budgets", tags=["Monthly Budgets"])


@router.post("/", response_model=MonthlyBudgetResponse, status_code=201)
async def get_or_create_budget(
    data: MonthlyBudgetCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    return await MonthlyBudgetService(db).get_or_create_budget(current_user.id, data)


@router.get("/by-month", response_model=MonthlyBudgetResponse)
async def get_budget_for_month(
    family_group_id: uuid.UUID,
    year: int,
    month: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    svc = MonthlyBudgetService(db)
    budget = await svc.get_budget_for_month(family_group_id, year, month)
    if budget is None:
        # Create it
        data = MonthlyBudgetCreate(family_group_id=family_group_id, year=year, month=month)
        return await svc.get_or_create_budget(current_user.id, data)
    return budget


@router.get("/{budget_id}", response_model=MonthlyBudgetResponse)
async def get_budget(
    budget_id: uuid.UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    return await MonthlyBudgetService(db).get_budget(budget_id)


@router.get("/{budget_id}/summary", response_model=BudgetSummaryResponse)
async def get_summary(
    budget_id: uuid.UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    return await MonthlyBudgetService(db).get_summary(budget_id)


@router.get("/{budget_id}/export")
async def export_budget(
    budget_id: uuid.UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Export monthly budget to Excel (.xlsx) file."""
    budget = await MonthlyBudgetService(db).get_budget(budget_id)

    wb = Workbook()
    ws = wb.active
    ws.title = f"Budżet {budget.month:02d}/{budget.year}"

    # Header styling
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill_income = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    header_fill_expense = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Title row
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = f"Budżet miesięczny — {budget.month:02d}/{budget.year}"
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    # Headers
    headers = ["Typ", "Nazwa", "Kwota (PLN)", "Cykliczne"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        cell.alignment = header_alignment
        cell.border = thin_border

    # Data rows
    row_idx = 4
    total_income = 0.0
    total_expenses = 0.0

    # Income section
    incomes = [e for e in budget.entries if e.entry_type.value == "income"]
    for entry in incomes:
        ws.cell(row=row_idx, column=1, value="Przychód").border = thin_border
        ws.cell(row=row_idx, column=2, value=entry.name).border = thin_border
        ws.cell(row=row_idx, column=3, value=float(entry.amount)).border = thin_border
        ws.cell(row=row_idx, column=3).number_format = "#,##0.00"
        ws.cell(row=row_idx, column=4, value="Tak" if entry.is_recurring else "Nie").border = thin_border
        total_income += float(entry.amount)
        row_idx += 1

    # Expenses section
    expenses = [e for e in budget.entries if e.entry_type.value == "expense"]
    for entry in expenses:
        ws.cell(row=row_idx, column=1, value="Wydatek").border = thin_border
        ws.cell(row=row_idx, column=2, value=entry.name).border = thin_border
        ws.cell(row=row_idx, column=3, value=float(entry.amount)).border = thin_border
        ws.cell(row=row_idx, column=3).number_format = "#,##0.00"
        ws.cell(row=row_idx, column=4, value="Tak" if entry.is_recurring else "Nie").border = thin_border
        total_expenses += float(entry.amount)
        row_idx += 1

    # Summary rows
    row_idx += 1
    summary_font = Font(bold=True, size=11)
    ws.cell(row=row_idx, column=1, value="SUMA PRZYCHODÓW").font = summary_font
    ws.cell(row=row_idx, column=1).border = thin_border
    ws.cell(row=row_idx, column=2).border = thin_border
    income_cell = ws.cell(row=row_idx, column=3, value=total_income)
    income_cell.font = Font(bold=True, color="059669")
    income_cell.number_format = "#,##0.00"
    income_cell.border = thin_border
    ws.cell(row=row_idx, column=4).border = thin_border
    row_idx += 1

    ws.cell(row=row_idx, column=1, value="SUMA WYDATKÓW").font = summary_font
    ws.cell(row=row_idx, column=1).border = thin_border
    ws.cell(row=row_idx, column=2).border = thin_border
    expense_cell = ws.cell(row=row_idx, column=3, value=total_expenses)
    expense_cell.font = Font(bold=True, color="DC2626")
    expense_cell.number_format = "#,##0.00"
    expense_cell.border = thin_border
    ws.cell(row=row_idx, column=4).border = thin_border
    row_idx += 1

    remaining = total_income - total_expenses
    ws.cell(row=row_idx, column=1, value="POZOSTAŁO").font = summary_font
    ws.cell(row=row_idx, column=1).border = thin_border
    ws.cell(row=row_idx, column=2).border = thin_border
    remaining_cell = ws.cell(row=row_idx, column=3, value=remaining)
    remaining_cell.font = Font(bold=True, color="059669" if remaining >= 0 else "DC2626")
    remaining_cell.number_format = "#,##0.00"
    remaining_cell.border = thin_border
    ws.cell(row=row_idx, column=4).border = thin_border

    # Column widths
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 14

    # Save to bytes
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"budzet_{budget.year}_{budget.month:02d}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/{budget_id}/entries", response_model=BudgetEntryResponse, status_code=201)
async def add_entry(
    budget_id: uuid.UUID,
    data: BudgetEntryCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    return await MonthlyBudgetService(db).add_entry(current_user.id, budget_id, data)


@router.patch("/entries/{entry_id}", response_model=BudgetEntryResponse)
async def update_entry(
    entry_id: uuid.UUID,
    data: BudgetEntryUpdate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    return await MonthlyBudgetService(db).update_entry_partial(entry_id, data)


@router.delete("/entries/{entry_id}", status_code=204)
async def remove_entry(
    entry_id: uuid.UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    await MonthlyBudgetService(db).remove_entry(entry_id)