import io
from datetime import date
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import get_current_user
from app.db.base import get_db
from app.models.user import User
from app.services.report import ReportService
from app.services.transaction import TransactionService

router = APIRouter(prefix="/reports", tags=["Reports"])


@router.get("/dashboard")
async def dashboard(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    return await ReportService(db).dashboard(current_user.id)


@router.get("/expenses-by-category")
async def expenses_by_category(
    year: int = Query(default=date.today().year),
    month: Optional[int] = None,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    return await ReportService(db).expenses_by_category(current_user.id, year, month)


@router.get("/monthly-trend")
async def monthly_trend(
    months: int = Query(default=12, ge=1, le=24),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    return await ReportService(db).monthly_trend(current_user.id, months)


@router.get("/export/transactions")
async def export_transactions(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    svc = TransactionService(db)
    transactions = await svc.list(
        current_user.id, start_date=start_date, end_date=end_date, limit=10_000
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transakcje"

    headers = ["Data", "Typ", "Kwota", "Waluta", "Kwota PLN", "Opis", "Zakres"]
    fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    bold_white = Font(color="FFFFFF", bold=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = fill
        cell.font = bold_white
        cell.alignment = Alignment(horizontal="center")

    for row, tx in enumerate(transactions, 2):
        ws.cell(row=row, column=1, value=str(tx.transaction_date))
        ws.cell(row=row, column=2, value=tx.transaction_type.value)
        ws.cell(row=row, column=3, value=float(tx.amount))
        ws.cell(row=row, column=4, value=tx.currency)
        ws.cell(row=row, column=5, value=float(tx.amount_pln) if tx.amount_pln else None)
        ws.cell(row=row, column=6, value=tx.description or "")
        ws.cell(row=row, column=7, value=tx.scope.value)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=transakcje.xlsx"},
    )
