import uuid
from datetime import date
from typing import Optional
from io import BytesIO

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.api.deps import get_current_user
from app.db.base import get_db
from app.models.user import User
from app.services.expense import SimpleExpenseService
from app.schemas.expense import SimpleExpenseCreate, SimpleExpenseResponse

router = APIRouter(prefix="/simple-expenses", tags=["Simple Expenses"])


@router.get("/", response_model=list[SimpleExpenseResponse])
async def list_expenses(
    family_group_id: uuid.UUID = Query(...),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    return await SimpleExpenseService(db).list_expenses(
        family_group_id, start_date, end_date
    )


@router.get("/export")
async def export_expenses(
    family_group_id: uuid.UUID = Query(...),
    year: int = Query(...),
    month: int = Query(...),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Export simple expenses to Excel (.xlsx) file."""
    import calendar

    start_date = date(year, month, 1)
    end_date = date(year, month, calendar.monthrange(year, month)[1])

    expenses = await SimpleExpenseService(db).list_expenses(
        family_group_id, start_date, end_date
    )

    wb = Workbook()
    ws = wb.active
    ws.title = f"Wydatki {month:02d}/{year}"

    # Header styling
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Title row
    ws.merge_cells("A1:E1")
    title_cell = ws["A1"]
    title_cell.value = f"Wydatki — {month:02d}/{year}"
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    # Headers
    headers = ["Data", "Opis", "Kwota (PLN)", "Użytkownik", "Data dodania"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Data rows
    total = 0.0
    for row_idx, expense in enumerate(expenses, 4):
        ws.cell(row=row_idx, column=1, value=str(expense.expense_date)).border = thin_border
        ws.cell(row=row_idx, column=2, value=expense.description).border = thin_border
        ws.cell(row=row_idx, column=3, value=float(expense.amount)).border = thin_border
        ws.cell(row=row_idx, column=3).number_format = "#,##0.00"
        ws.cell(row=row_idx, column=4, value=str(expense.user_id)).border = thin_border
        ws.cell(row=row_idx, column=5, value=str(expense.created_at)).border = thin_border
        total += float(expense.amount)

    # Total row
    total_row = len(expenses) + 4
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=2)
    total_label = ws.cell(row=total_row, column=1, value="Suma")
    total_label.font = Font(bold=True)
    total_label.alignment = Alignment(horizontal="right")
    total_label.border = thin_border
    ws.cell(row=total_row, column=1).border = thin_border

    total_cell = ws.cell(row=total_row, column=3, value=total)
    total_cell.font = Font(bold=True)
    total_cell.number_format = "#,##0.00"
    total_cell.border = thin_border

    # Column widths
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 40
    ws.column_dimensions["E"].width = 22

    # Save to bytes
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"wydatki_{year}_{month:02d}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/", response_model=SimpleExpenseResponse, status_code=201)
async def create_expense(
    data: SimpleExpenseCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    return await SimpleExpenseService(db).create_expense(current_user.id, data)


@router.delete("/{expense_id}", status_code=204)
async def delete_expense(
    expense_id: uuid.UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    await SimpleExpenseService(db).delete_expense(current_user.id, expense_id)